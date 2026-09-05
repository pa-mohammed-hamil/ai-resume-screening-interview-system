# Project scaffold file
"""
genai/rag/document_loader.py

Document loading utilities for the RAG pipeline.

Supported formats:
    - PDF
    - DOCX
    - TXT
    - MD

The loader converts source files into a common Document structure that can
be passed to the chunker -> embedding -> vector store pipeline.
"""

from __future__ import annotations

import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable

logger = logging.getLogger(__name__)


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt",
    ".md",
    ".markdown",
}


@dataclass
class Document:
    """Normalized document representation used by the RAG pipeline."""

    content: str
    metadata: dict[str, Any] = field(default_factory=dict)

    @property
    def text(self) -> str:
        """Backward-compatible alias for document content."""
        return self.content


class DocumentLoadError(Exception):
    """Raised when a document cannot be loaded."""


class UnsupportedDocumentError(DocumentLoadError):
    """Raised when the file type is not supported."""


class DocumentLoader:
    """
    Load documents from files or directories.

    Example:
        loader = DocumentLoader()

        documents = loader.load("data/raw/resumes/resume.pdf")

        for document in documents:
            print(document.content)
    """

    def __init__(
        self,
        supported_extensions: Iterable[str] | None = None,
        encoding: str = "utf-8",
    ) -> None:
        self.supported_extensions = {
            ext.lower() if ext.startswith(".") else f".{ext.lower()}"
            for ext in (
                supported_extensions
                if supported_extensions is not None
                else SUPPORTED_EXTENSIONS
            )
        }

        self.encoding = encoding

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------

    def load(self, source: str | Path) -> list[Document]:
        """
        Load a single file or every supported file inside a directory.

        Args:
            source: File or directory path.

        Returns:
            List of normalized Document objects.

        Raises:
            FileNotFoundError: If the source does not exist.
            UnsupportedDocumentError: If a file format is unsupported.
            DocumentLoadError: If loading fails.
        """

        path = Path(source)

        if not path.exists():
            raise FileNotFoundError(f"Document source not found: {path}")

        if path.is_dir():
            return self.load_directory(path)

        return [self.load_file(path)]

    def load_file(self, path: str | Path) -> Document:
        """
        Load one document.

        Args:
            path: Path to the document.

        Returns:
            Document object.
        """

        path = Path(path)

        if not path.exists():
            raise FileNotFoundError(f"File not found: {path}")

        if not path.is_file():
            raise DocumentLoadError(f"Not a file: {path}")

        extension = path.suffix.lower()

        if extension not in self.supported_extensions:
            raise UnsupportedDocumentError(
                f"Unsupported document type: {extension}. "
                f"Supported types: {sorted(self.supported_extensions)}"
            )

        try:
            if extension == ".pdf":
                content = self._load_pdf(path)

            elif extension == ".docx":
                content = self._load_docx(path)

            elif extension in {".txt", ".md", ".markdown"}:
                content = self._load_text(path)

            else:
                raise UnsupportedDocumentError(
                    f"Unsupported extension: {extension}"
                )

        except DocumentLoadError:
            raise

        except Exception as exc:
            logger.exception("Failed to load document: %s", path)
            raise DocumentLoadError(
                f"Failed to load document '{path}': {exc}"
            ) from exc

        content = self._clean_text(content)

        if not content:
            raise DocumentLoadError(
                f"Document contains no readable text: {path}"
            )

        metadata = self._build_metadata(path)

        return Document(
            content=content,
            metadata=metadata,
        )

    def load_directory(
        self,
        directory: str | Path,
        recursive: bool = True,
    ) -> list[Document]:
        """
        Load all supported documents from a directory.

        Args:
            directory: Directory containing documents.
            recursive: Search subdirectories when True.

        Returns:
            List of loaded documents.
        """

        directory = Path(directory)

        if not directory.exists():
            raise FileNotFoundError(
                f"Directory not found: {directory}"
            )

        if not directory.is_dir():
            raise DocumentLoadError(
                f"Expected directory but received: {directory}"
            )

        pattern = "**/*" if recursive else "*"

        documents: list[Document] = []

        for path in sorted(directory.glob(pattern)):
            if not path.is_file():
                continue

            if path.suffix.lower() not in self.supported_extensions:
                continue

            try:
                documents.append(self.load_file(path))

            except DocumentLoadError as exc:
                logger.warning(
                    "Skipping document '%s': %s",
                    path,
                    exc,
                )

        return documents

    # ------------------------------------------------------------------
    # Format-specific loaders
    # ------------------------------------------------------------------

    def _load_pdf(self, path: Path) -> str:
        """Extract text from a PDF document."""

        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise DocumentLoadError(
                "pypdf is required for PDF processing. "
                "Install it with: pip install pypdf"
            ) from exc

        reader = PdfReader(str(path))

        pages: list[str] = []

        for page_number, page in enumerate(reader.pages, start=1):
            try:
                text = page.extract_text() or ""

                if text.strip():
                    pages.append(
                        f"[Page {page_number}]\n{text}"
                    )

            except Exception as exc:
                logger.warning(
                    "Could not extract page %s from %s: %s",
                    page_number,
                    path,
                    exc,
                )

        return "\n\n".join(pages)

    def _load_docx(self, path: Path) -> str:
        """Extract paragraphs and tables from a DOCX document."""

        try:
            from docx import Document as DocxDocument
        except ImportError as exc:
            raise DocumentLoadError(
                "python-docx is required for DOCX processing. "
                "Install it with: pip install python-docx"
            ) from exc

        document = DocxDocument(str(path))

        parts: list[str] = []

        # Paragraphs
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                parts.append(text)

        # Tables
        for table_index, table in enumerate(
            document.tables,
            start=1,
        ):
            rows: list[str] = []

            for row in table.rows:
                cells = [
                    cell.text.strip()
                    for cell in row.cells
                ]

                cells = [cell for cell in cells if cell]

                if cells:
                    rows.append(" | ".join(cells))

            if rows:
                parts.append(
                    f"[Table {table_index}]\n"
                    + "\n".join(rows)
                )

        return "\n\n".join(parts)

    def _load_text(self, path: Path) -> str:
        """Load TXT, Markdown, or similar text files."""

        try:
            return path.read_text(
                encoding=self.encoding,
                errors="replace",
            )

        except UnicodeDecodeError:
            logger.warning(
                "UTF-8 decoding failed for %s. "
                "Trying latin-1.",
                path,
            )

            return path.read_text(
                encoding="latin-1",
                errors="replace",
            )

    # ------------------------------------------------------------------
    # Metadata
    # ------------------------------------------------------------------

    def _build_metadata(self, path: Path) -> dict[str, Any]:
        """Create metadata used by the RAG pipeline."""

        try:
            stat = path.stat()

            file_size = stat.st_size
            modified_time = stat.st_mtime

        except OSError:
            file_size = None
            modified_time = None

        return {
            "source": str(path.resolve()),
            "file_name": path.name,
            "file_stem": path.stem,
            "file_type": path.suffix.lower(),
            "file_size": file_size,
            "modified_time": modified_time,
        }

    # ------------------------------------------------------------------
    # Text normalization
    # ------------------------------------------------------------------

    @staticmethod
    def _clean_text(text: str) -> str:
        """
        Normalize extracted text without destroying useful structure.
        """

        if not text:
            return ""

        # Normalize line endings.
        text = text.replace("\r\n", "\n")
        text = text.replace("\r", "\n")

        # Remove null bytes.
        text = text.replace("\x00", "")

        # Remove trailing whitespace from each line.
        lines = [
            line.rstrip()
            for line in text.split("\n")
        ]

        # Collapse excessive blank lines.
        cleaned_lines: list[str] = []

        previous_blank = False

        for line in lines:
            is_blank = not line.strip()

            if is_blank:
                if previous_blank:
                    continue

                cleaned_lines.append("")
                previous_blank = True

            else:
                cleaned_lines.append(line)
                previous_blank = False

        return "\n".join(cleaned_lines).strip()


# ----------------------------------------------------------------------
# Convenience functions
# ----------------------------------------------------------------------

_default_loader = DocumentLoader()


def load_document(source: str | Path) -> list[Document]:
    """
    Convenience function for loading a file or directory.

    Example:
        documents = load_document("data/raw/resumes")
    """
    return _default_loader.load(source)


def load_file(path: str | Path) -> Document:
    """
    Convenience function for loading one file.
    """
    return _default_loader.load_file(path)


def load_directory(
    directory: str | Path,
    recursive: bool = True,
) -> list[Document]:
    """
    Convenience function for loading a directory.
    """
    return _default_loader.load_directory(
        directory,
        recursive=recursive,
    )


__all__ = [
    "Document",
    "DocumentLoader",
    "DocumentLoadError",
    "UnsupportedDocumentError",
    "load_document",
    "load_file",
    "load_directory",
]