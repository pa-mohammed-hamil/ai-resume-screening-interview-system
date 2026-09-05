# Project scaffold file
"""
AI Resume Screening & Interview System
=======================================

File:
    scripts/ingest_documents.py

Purpose:
    Ingest raw resumes and job descriptions, clean their
    contents, extract basic metadata, and save normalized
    documents into data/processed/.

Input:
    data/raw/resumes/
    data/raw/job_descriptions/

Output:
    data/processed/resumes/
    data/processed/job_descriptions/

Supported files:
    .txt
    .md
    .json
    .pdf
    .docx

Usage:
    python scripts/ingest_documents.py

    python scripts/ingest_documents.py --type resumes

    python scripts/ingest_documents.py --type jobs

    python scripts/ingest_documents.py --type all

    python scripts/ingest_documents.py --force
"""

from __future__ import annotations

import argparse
import hashlib
import json
import logging
import re
import shutil
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any


# ============================================================
# PROJECT PATHS
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parents[1]

DATA_DIR = PROJECT_ROOT / "data"

RAW_DIR = DATA_DIR / "raw"

PROCESSED_DIR = DATA_DIR / "processed"

RAW_RESUMES_DIR = RAW_DIR / "resumes"

RAW_JOBS_DIR = RAW_DIR / "job_descriptions"

PROCESSED_RESUMES_DIR = (
    PROCESSED_DIR / "resumes"
)

PROCESSED_JOBS_DIR = (
    PROCESSED_DIR / "job_descriptions"
)

MANIFEST_PATH = (
    PROCESSED_DIR / "ingestion_manifest.json"
)


# ============================================================
# LOGGING
# ============================================================

logging.basicConfig(
    level=logging.INFO,
    format=(
        "%(asctime)s | "
        "%(levelname)s | "
        "%(message)s"
    ),
)

logger = logging.getLogger(
    "document-ingestion"
)


# ============================================================
# SUPPORTED FILE TYPES
# ============================================================

SUPPORTED_TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".json",
}

SUPPORTED_DOCUMENT_EXTENSIONS = {
    ".pdf",
    ".docx",
}

SUPPORTED_EXTENSIONS = (
    SUPPORTED_TEXT_EXTENSIONS
    | SUPPORTED_DOCUMENT_EXTENSIONS
)


# ============================================================
# DIRECTORY SETUP
# ============================================================

def create_directories() -> None:
    """
    Create all required ingestion directories.
    """

    directories = [
        RAW_RESUMES_DIR,
        RAW_JOBS_DIR,
        PROCESSED_RESUMES_DIR,
        PROCESSED_JOBS_DIR,
    ]

    for directory in directories:

        directory.mkdir(
            parents=True,
            exist_ok=True,
        )


# ============================================================
# TEXT CLEANING
# ============================================================

def normalize_unicode(
    text: str,
) -> str:
    """
    Normalize common Unicode characters.
    """

    replacements = {
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2013": "-",
        "\u2014": "-",
        "\u2212": "-",
        "\u00a0": " ",
        "\u2022": "-",
        "\u2026": "...",
    }

    for old, new in replacements.items():

        text = text.replace(
            old,
            new,
        )

    return text


def clean_text(
    text: str,
) -> str:
    """
    Clean and normalize document text.

    Preserves useful line structure while removing
    excessive whitespace.
    """

    if not text:

        return ""

    text = normalize_unicode(
        text
    )

    # Remove null bytes.
    text = text.replace(
        "\x00",
        "",
    )

    # Normalize Windows/Mac line endings.
    text = text.replace(
        "\r\n",
        "\n",
    ).replace(
        "\r",
        "\n",
    )

    # Remove excessive spaces/tabs.
    text = re.sub(
        r"[ \t]+",
        " ",
        text,
    )

    # Remove excessive blank lines.
    text = re.sub(
        r"\n{3,}",
        "\n\n",
        text,
    )

    # Clean individual lines.
    lines = []

    for line in text.splitlines():

        line = line.strip()

        if line:

            lines.append(
                line
            )

    return "\n".join(
        lines
    ).strip()


# ============================================================
# HASHING
# ============================================================

def calculate_hash(
    content: bytes,
) -> str:
    """
    Calculate SHA-256 hash for raw content.
    """

    return hashlib.sha256(
        content
    ).hexdigest()


def calculate_text_hash(
    text: str,
) -> str:
    """
    Calculate SHA-256 hash for normalized text.
    """

    return hashlib.sha256(
        text.encode(
            "utf-8"
        )
    ).hexdigest()


# ============================================================
# FILE READING
# ============================================================

def read_text_file(
    path: Path,
) -> str:
    """
    Read TXT or Markdown file.
    """

    try:

        return path.read_text(
            encoding="utf-8",
            errors="ignore",
        )

    except OSError as exc:

        logger.error(
            "Unable to read %s: %s",
            path,
            exc,
        )

        return ""


def read_json_file(
    path: Path,
) -> str:
    """
    Read JSON and convert it to normalized text.
    """

    try:

        content = path.read_text(
            encoding="utf-8",
            errors="ignore",
        )

        data = json.loads(
            content
        )

        if isinstance(
            data,
            dict,
        ):

            # Preserve useful fields in readable form.
            parts = []

            for key, value in data.items():

                if isinstance(
                    value,
                    (dict, list),
                ):

                    value = json.dumps(
                        value,
                        ensure_ascii=False,
                    )

                parts.append(
                    f"{key}: {value}"
                )

            return "\n".join(
                parts
            )

        return json.dumps(
            data,
            ensure_ascii=False,
        )

    except (
        OSError,
        json.JSONDecodeError,
    ) as exc:

        logger.error(
            "Unable to parse JSON %s: %s",
            path,
            exc,
        )

        return ""


def read_pdf_file(
    path: Path,
) -> str:
    """
    Extract text from PDF using pypdf.
    """

    try:

        from pypdf import (
            PdfReader,
        )

    except ImportError:

        logger.error(
            "pypdf is not installed. "
            "Install it with: pip install pypdf"
        )

        return ""

    try:

        reader = PdfReader(
            str(path)
        )

        pages = []

        for page in reader.pages:

            try:

                text = page.extract_text()

            except Exception as exc:

                logger.warning(
                    "Failed to extract PDF page "
                    "from %s: %s",
                    path,
                    exc,
                )

                continue

            if text:

                pages.append(
                    text
                )

        return "\n\n".join(
            pages
        )

    except Exception as exc:

        logger.error(
            "Unable to read PDF %s: %s",
            path,
            exc,
        )

        return ""


def read_docx_file(
    path: Path,
) -> str:
    """
    Extract text from DOCX using python-docx.
    """

    try:

        from docx import (
            Document,
        )

    except ImportError:

        logger.error(
            "python-docx is not installed. "
            "Install it with: pip install python-docx"
        )

        return ""

    try:

        document = Document(
            str(path)
        )

        paragraphs = []

        for paragraph in document.paragraphs:

            text = paragraph.text.strip()

            if text:

                paragraphs.append(
                    text
                )

        # Also collect table contents.
        for table in document.tables:

            for row in table.rows:

                cells = []

                for cell in row.cells:

                    value = cell.text.strip()

                    if value:

                        cells.append(
                            value
                        )

                if cells:

                    paragraphs.append(
                        " | ".join(
                            cells
                        )
                    )

        return "\n".join(
            paragraphs
        )

    except Exception as exc:

        logger.error(
            "Unable to read DOCX %s: %s",
            path,
            exc,
        )

        return ""


def extract_text(
    path: Path,
) -> str:
    """
    Extract text from a supported document.
    """

    extension = (
        path.suffix.lower()
    )

    if extension in (
        ".txt",
        ".md",
    ):

        return read_text_file(
            path
        )

    if extension == ".json":

        return read_json_file(
            path
        )

    if extension == ".pdf":

        return read_pdf_file(
            path
        )

    if extension == ".docx":

        return read_docx_file(
            path
        )

    logger.warning(
        "Unsupported file type: %s",
        path,
    )

    return ""


# ============================================================
# DOCUMENT DISCOVERY
# ============================================================

def discover_documents(
    directory: Path,
) -> list[Path]:
    """
    Find supported documents recursively.
    """

    if not directory.exists():

        logger.warning(
            "Directory does not exist: %s",
            directory,
        )

        return []

    documents = []

    for path in directory.rglob("*"):

        if not path.is_file():
            continue

        if (
            path.suffix.lower()
            not in SUPPORTED_EXTENSIONS
        ):

            continue

        documents.append(
            path
        )

    return sorted(
        documents
    )


# ============================================================
# FILE NAME SANITIZATION
# ============================================================

def sanitize_filename(
    filename: str,
) -> str:
    """
    Create a safe filename for processed documents.
    """

    stem = Path(
        filename
    ).stem

    stem = stem.lower()

    stem = re.sub(
        r"[^a-z0-9_-]+",
        "_",
        stem,
    )

    stem = re.sub(
        r"_+",
        "_",
        stem,
    )

    stem = stem.strip(
        "_"
    )

    if not stem:

        stem = "document"

    return stem


def build_output_filename(
    source: Path,
    source_root: Path,
) -> str:
    """
    Generate a collision-resistant output filename.
    """

    relative = source.relative_to(
        source_root
    )

    safe_stem = sanitize_filename(
        source.name
    )

    relative_hash = hashlib.sha1(
        str(
            relative
        ).encode(
            "utf-8"
        )
    ).hexdigest()[:10]

    return (
        f"{safe_stem}_{relative_hash}.txt"
    )


# ============================================================
# METADATA
# ============================================================

def detect_document_type(
    category: str,
) -> str:

    if category == "resumes":

        return "resume"

    if category == "jobs":

        return "job_description"

    return "document"


def extract_basic_metadata(
    text: str,
    document_type: str,
) -> dict[str, Any]:
    """
    Extract lightweight metadata.

    More advanced extraction is handled by the AI modules
    under ai/information_extraction/.
    """

    words = re.findall(
        r"\b[\w+#.-]+\b",
        text,
    )

    lines = [
        line.strip()
        for line in text.splitlines()
        if line.strip()
    ]

    metadata: dict[str, Any] = {
        "document_type": document_type,
        "character_count": len(text),
        "word_count": len(words),
        "line_count": len(lines),
    }

    # --------------------------------------------------------
    # Email detection
    # --------------------------------------------------------

    email_match = re.search(
        r"\b[A-Za-z0-9._%+-]+"
        r"@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b",
        text,
    )

    if email_match:

        metadata[
            "email"
        ] = email_match.group(
            0
        )

    # --------------------------------------------------------
    # Phone detection
    # --------------------------------------------------------

    phone_match = re.search(
        r"(?<!\d)"
        r"(?:\+?\d{1,3}[\s.-]?)?"
        r"(?:\(?\d{3}\)?[\s.-]?)?"
        r"\d{3}[\s.-]\d{4}"
        r"(?!\d)",
        text,
    )

    if phone_match:

        metadata[
            "phone"
        ] = phone_match.group(
            0
        )

    # --------------------------------------------------------
    # Common resume sections
    # --------------------------------------------------------

    if document_type == "resume":

        sections = {
            "education": [
                "education",
                "academic",
                "qualification",
            ],
            "experience": [
                "experience",
                "employment",
                "work history",
            ],
            "skills": [
                "skills",
                "technical skills",
                "competencies",
            ],
            "projects": [
                "projects",
                "personal projects",
            ],
            "certifications": [
                "certifications",
                "certificates",
            ],
        }

        lower_text = text.lower()

        detected_sections = []

        for section, keywords in sections.items():

            if any(
                keyword in lower_text
                for keyword in keywords
            ):

                detected_sections.append(
                    section
                )

        metadata[
            "detected_sections"
        ] = detected_sections

    # --------------------------------------------------------
    # Job description sections
    # --------------------------------------------------------

    if document_type == "job_description":

        sections = {
            "requirements": [
                "requirements",
                "qualifications",
                "required skills",
            ],
            "responsibilities": [
                "responsibilities",
                "what you'll do",
                "duties",
            ],
            "benefits": [
                "benefits",
                "perks",
            ],
        }

        lower_text = text.lower()

        detected_sections = []

        for section, keywords in sections.items():

            if any(
                keyword in lower_text
                for keyword in keywords
            ):

                detected_sections.append(
                    section
                )

        metadata[
            "detected_sections"
        ] = detected_sections

    return metadata


# ============================================================
# PROCESSED DOCUMENT FORMAT
# ============================================================

def build_processed_document(
    source: Path,
    text: str,
    category: str,
) -> dict[str, Any]:

    document_type = detect_document_type(
        category
    )

    raw_bytes = source.read_bytes()

    normalized_text = clean_text(
        text
    )

    return {
        "document_id": create_document_id(
            source
        ),
        "document_type": document_type,
        "source_filename": source.name,
        "source_extension": source.suffix.lower(),
        "source_path": str(
            source.relative_to(
                PROJECT_ROOT
            )
        ),
        "raw_hash": calculate_hash(
            raw_bytes
        ),
        "text_hash": calculate_text_hash(
            normalized_text
        ),
        "ingested_at": datetime.now(
            timezone.utc
        ).isoformat(),
        "metadata": extract_basic_metadata(
            normalized_text,
            document_type,
        ),
        "text": normalized_text,
    }


def create_document_id(
    source: Path,
) -> str:

    relative = str(
        source.relative_to(
            PROJECT_ROOT
        )
    )

    return hashlib.sha256(
        relative.encode(
            "utf-8"
        )
    ).hexdigest()[:24]


# ============================================================
# WRITE PROCESSED FILE
# ============================================================

def write_processed_text(
    path: Path,
    document: dict[str, Any],
) -> None:
    """
    Write normalized text.

    The processed text is intentionally stored as plain text
    so it can be consumed directly by the embedding and RAG
    pipeline.
    """

    path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    path.write_text(
        document["text"],
        encoding="utf-8",
    )


def write_metadata(
    path: Path,
    document: dict[str, Any],
) -> None:
    """
    Write metadata alongside processed document.
    """

    metadata_path = path.with_suffix(
        ".json"
    )

    metadata = {
        key: value
        for key, value in document.items()
        if key != "text"
    }

    metadata_path.write_text(
        json.dumps(
            metadata,
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )


# ============================================================
# MANIFEST
# ============================================================

def load_manifest() -> dict[str, Any]:

    if not MANIFEST_PATH.exists():

        return {
            "version": "1.0",
            "documents": [],
        }

    try:

        data = json.loads(
            MANIFEST_PATH.read_text(
                encoding="utf-8"
            )
        )

        if not isinstance(
            data,
            dict,
        ):

            raise ValueError(
                "Invalid manifest."
            )

        if not isinstance(
            data.get(
                "documents"
            ),
            list,
        ):

            data[
                "documents"
            ] = []

        return data

    except (
        OSError,
        json.JSONDecodeError,
        ValueError,
    ) as exc:

        logger.warning(
            "Unable to load ingestion manifest: %s",
            exc,
        )

        return {
            "version": "1.0",
            "documents": [],
        }


def save_manifest(
    manifest: dict[str, Any],
) -> None:

    MANIFEST_PATH.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    temporary_path = (
        MANIFEST_PATH.with_suffix(
            ".tmp"
        )
    )

    temporary_path.write_text(
        json.dumps(
            manifest,
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    temporary_path.replace(
        MANIFEST_PATH
    )


# ============================================================
# INGEST ONE DOCUMENT
# ============================================================

def ingest_document(
    source: Path,
    source_root: Path,
    output_root: Path,
    category: str,
    force: bool,
) -> dict[str, Any] | None:

    logger.info(
        "Processing: %s",
        source,
    )

    # --------------------------------------------------------
    # Read
    # --------------------------------------------------------

    text = extract_text(
        source
    )

    if not text.strip():

        logger.warning(
            "No text extracted from %s",
            source,
        )

        return None

    # --------------------------------------------------------
    # Build document
    # --------------------------------------------------------

    document = build_processed_document(
        source=source,
        text=text,
        category=category,
    )

    # --------------------------------------------------------
    # Output filename
    # --------------------------------------------------------

    filename = build_output_filename(
        source,
        source_root,
    )

    output_path = (
        output_root / filename
    )

    # --------------------------------------------------------
    # Skip unchanged document
    # --------------------------------------------------------

    if output_path.exists() and not force:

        metadata_path = (
            output_path.with_suffix(
                ".json"
            )
        )

        if metadata_path.exists():

            try:

                existing = json.loads(
                    metadata_path.read_text(
                        encoding="utf-8"
                    )
                )

                if (
                    existing.get(
                        "raw_hash"
                    )
                    == document.get(
                        "raw_hash"
                    )
                ):

                    logger.info(
                        "Unchanged, skipping: %s",
                        source,
                    )

                    document[
                        "processed_path"
                    ] = str(
                        output_path.relative_to(
                            PROJECT_ROOT
                        )
                    )

                    return document

            except (
                OSError,
                json.JSONDecodeError,
            ):
                pass

    # --------------------------------------------------------
    # Write files
    # --------------------------------------------------------

    write_processed_text(
        output_path,
        document,
    )

    write_metadata(
        output_path,
        document,
    )

    document[
        "processed_path"
    ] = str(
        output_path.relative_to(
            PROJECT_ROOT
        )
    )

    logger.info(
        "Created: %s",
        output_path,
    )

    return document


# ============================================================
# INGEST CATEGORY
# ============================================================

def ingest_category(
    category: str,
    force: bool,
    manifest: dict[str, Any],
) -> tuple[int, int]:

    if category == "resumes":

        source_root = RAW_RESUMES_DIR

        output_root = (
            PROCESSED_RESUMES_DIR
        )

    elif category == "jobs":

        source_root = RAW_JOBS_DIR

        output_root = (
            PROCESSED_JOBS_DIR
        )

    else:

        raise ValueError(
            f"Unsupported category: {category}"
        )

    documents = discover_documents(
        source_root
    )

    logger.info(
        "Found %d %s.",
        len(documents),
        category,
    )

    processed = 0

    failed = 0

    manifest_documents = {
        item.get(
            "source_path"
        ): item
        for item in manifest.get(
            "documents",
            []
        )
        if isinstance(
            item,
            dict,
        )
        and item.get(
            "source_path"
        )
    }

    for source in documents:

        try:

            document = ingest_document(
                source=source,
                source_root=source_root,
                output_root=output_root,
                category=category,
                force=force,
            )

            if document is None:

                failed += 1

                continue

            manifest_documents[
                document[
                    "source_path"
                ]
            ] = document

            processed += 1

        except Exception as exc:

            failed += 1

            logger.exception(
                "Failed to process %s: %s",
                source,
                exc,
            )

    manifest[
        "documents"
    ] = sorted(
        manifest_documents.values(),
        key=lambda item: item.get(
            "source_path",
            "",
        ),
    )

    return (
        processed,
        failed,
    )


# ============================================================
# REMOVE OLD OUTPUT
# ============================================================

def remove_processed_files(
    directory: Path,
) -> None:
    """
    Remove generated processed files.

    Used by --force-clean.
    """

    if not directory.exists():

        return

    for path in directory.iterdir():

        if path.is_file():

            path.unlink()

        elif path.is_dir():

            shutil.rmtree(
                path
            )


# ============================================================
# INGESTION PIPELINE
# ============================================================

def run_ingestion(
    document_type: str,
    force: bool,
    force_clean: bool,
) -> dict[str, Any]:

    create_directories()

    manifest = load_manifest()

    total_processed = 0

    total_failed = 0

    # --------------------------------------------------------
    # Clean requested category
    # --------------------------------------------------------

    if force_clean:

        if document_type in (
            "resumes",
            "all",
        ):

            remove_processed_files(
                PROCESSED_RESUMES_DIR
            )

        if document_type in (
            "jobs",
            "all",
        ):

            remove_processed_files(
                PROCESSED_JOBS_DIR
            )

    # --------------------------------------------------------
    # Resume ingestion
    # --------------------------------------------------------

    if document_type in (
        "resumes",
        "all",
    ):

        (
            processed,
            failed,
        ) = ingest_category(
            category="resumes",
            force=force,
            manifest=manifest,
        )

        total_processed += processed

        total_failed += failed

    # --------------------------------------------------------
    # Job ingestion
    # --------------------------------------------------------

    if document_type in (
        "jobs",
        "all",
    ):

        (
            processed,
            failed,
        ) = ingest_category(
            category="jobs",
            force=force,
            manifest=manifest,
        )

        total_processed += processed

        total_failed += failed

    # --------------------------------------------------------
    # Manifest metadata
    # --------------------------------------------------------

    manifest[
        "version"
    ] = "1.0"

    manifest[
        "updated_at"
    ] = datetime.now(
        timezone.utc
    ).isoformat()

    manifest[
        "statistics"
    ] = {
        "total_documents": len(
            manifest[
                "documents"
            ]
        ),
        "processed_this_run": (
            total_processed
        ),
        "failed_this_run": (
            total_failed
        ),
    }

    save_manifest(
        manifest
    )

    return {
        "processed": total_processed,
        "failed": total_failed,
        "total": len(
            manifest[
                "documents"
            ]
        ),
    }


# ============================================================
# REPORT
# ============================================================

def print_report(
    result: dict[str, Any],
) -> None:

    print()
    print("=" * 72)
    print(
        "AI RESUME SCREENING & INTERVIEW SYSTEM"
    )
    print(
        "DOCUMENT INGESTION REPORT"
    )
    print("=" * 72)

    print(
        f"Processed this run : "
        f"{result['processed']}"
    )

    print(
        f"Failed             : "
        f"{result['failed']}"
    )

    print(
        f"Manifest documents  : "
        f"{result['total']}"
    )

    print(
        f"Manifest            : "
        f"{MANIFEST_PATH}"
    )

    print()
    print(
        "Processed resumes:"
    )

    print(
        f"  {PROCESSED_RESUMES_DIR}"
    )

    print(
        "Processed jobs:"
    )

    print(
        f"  {PROCESSED_JOBS_DIR}"
    )

    print("=" * 72)
    print(
        "Document ingestion completed."
    )
    print("=" * 72)
    print()


# ============================================================
# ARGUMENTS
# ============================================================

def parse_arguments() -> argparse.Namespace:

    parser = argparse.ArgumentParser(
        description=(
            "Ingest and normalize resumes "
            "and job descriptions."
        )
    )

    parser.add_argument(
        "--type",
        choices=[
            "resumes",
            "jobs",
            "all",
        ],
        default="all",
        help=(
            "Document category to ingest."
        ),
    )

    parser.add_argument(
        "--force",
        action="store_true",
        help=(
            "Reprocess documents even when "
            "their content has not changed."
        ),
    )

    parser.add_argument(
        "--force-clean",
        action="store_true",
        help=(
            "Delete processed files for the "
            "selected category before ingestion."
        ),
    )

    parser.add_argument(
        "--quiet",
        action="store_true",
        help=(
            "Suppress the final report."
        ),
    )

    return parser.parse_args()


# ============================================================
# VALIDATION
# ============================================================

def validate_environment() -> None:

    create_directories()

    if not RAW_DIR.exists():

        raise FileNotFoundError(
            f"Raw data directory does not exist: "
            f"{RAW_DIR}"
        )


# ============================================================
# MAIN
# ============================================================

def main() -> int:

    args = parse_arguments()

    try:

        validate_environment()

        logger.info(
            "Starting document ingestion."
        )

        logger.info(
            "Document type: %s",
            args.type,
        )

        result = run_ingestion(
            document_type=args.type,
            force=args.force,
            force_clean=args.force_clean,
        )

        if not args.quiet:

            print_report(
                result
            )

        if result[
            "failed"
        ] > 0:

            return 1

        return 0

    except KeyboardInterrupt:

        logger.warning(
            "Document ingestion cancelled."
        )

        return 130

    except Exception as exc:

        logger.exception(
            "Document ingestion failed: %s",
            exc,
        )

        return 1


# ============================================================
# ENTRY POINT
# ============================================================

if __name__ == "__main__":

    sys.exit(
        main()
    )