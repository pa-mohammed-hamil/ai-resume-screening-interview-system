# Project scaffold file
"""
DOCX Resume Parser
------------------
Extracts text and metadata from DOCX resumes.

Used by:
    ai/resume_parser/parser.py
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any


class DOCXParserError(Exception):
    """Raised when DOCX parsing fails."""


class DOCXParser:
    """Parser for Microsoft Word DOCX resume documents."""

    def __init__(
        self,
        include_tables: bool = True,
        include_headers_footers: bool = True,
        include_hyperlinks: bool = True,
    ) -> None:
        self.include_tables = include_tables
        self.include_headers_footers = include_headers_footers
        self.include_hyperlinks = include_hyperlinks

    def parse(
        self,
        file_path: str | Path,
    ) -> tuple[str, dict[str, Any]]:
        """
        Extract text from a DOCX file.

        Args:
            file_path: Path to DOCX resume.

        Returns:
            Tuple containing extracted text and metadata.
        """

        path = Path(file_path)

        if not path.exists():
            raise FileNotFoundError(
                f"DOCX file not found: {path}"
            )

        if not path.is_file():
            raise DOCXParserError(
                f"Expected a file: {path}"
            )

        if path.suffix.lower() != ".docx":
            raise DOCXParserError(
                f"Expected a DOCX file, got: {path.suffix}"
            )

        try:
            file_bytes = path.read_bytes()

            return self.parse_bytes(
                file_bytes,
                path.name,
            )

        except DOCXParserError:
            raise

        except Exception as exc:
            raise DOCXParserError(
                f"Failed to read DOCX '{path.name}': {exc}"
            ) from exc

    def parse_bytes(
        self,
        file_bytes: bytes,
        file_name: str = "resume.docx",
    ) -> tuple[str, dict[str, Any]]:
        """
        Extract text from DOCX bytes.

        Useful for FastAPI UploadFile.

        Args:
            file_bytes: DOCX contents.
            file_name: Original filename.

        Returns:
            Tuple containing extracted text and metadata.
        """

        if not file_bytes:
            raise DOCXParserError(
                "DOCX file is empty."
            )

        if not self._looks_like_docx(file_bytes):
            raise DOCXParserError(
                f"'{file_name}' does not appear to be a valid DOCX file."
            )

        try:
            from docx import Document
        except ImportError as exc:
            raise DOCXParserError(
                "python-docx is required for DOCX parsing. "
                "Install it with: pip install python-docx"
            ) from exc

        try:
            document = Document(
                BytesIO(file_bytes)
            )
        except Exception as exc:
            raise DOCXParserError(
                f"Unable to open DOCX '{file_name}': {exc}"
            ) from exc

        try:
            return self._extract_document(
                document,
                file_name,
            )
        except DOCXParserError:
            raise
        except Exception as exc:
            raise DOCXParserError(
                f"Failed to extract DOCX content: {exc}"
            ) from exc

    def _extract_document(
        self,
        document: Any,
        file_name: str,
    ) -> tuple[str, dict[str, Any]]:
        """Extract paragraphs, tables, headers, and footers."""

        sections: list[str] = []

        paragraph_count = 0
        table_count = 0
        table_row_count = 0

        # Main document paragraphs
        for paragraph in document.paragraphs:
            text = self._extract_paragraph_text(
                paragraph
            )

            if text:
                sections.append(text)
                paragraph_count += 1

        # Tables
        if self.include_tables:
            for table_index, table in enumerate(
                document.tables,
                start=1,
            ):
                table_text, rows = self._extract_table(
                    table
                )

                if table_text:
                    sections.append(
                        f"[Table {table_index}]\n"
                        f"{table_text}"
                    )

                table_count += 1
                table_row_count += rows

        # Headers and footers
        header_footer_count = 0

        if self.include_headers_footers:
            for section in document.sections:

                header_text = self._extract_container(
                    section.header
                )

                if header_text:
                    sections.insert(
                        0,
                        header_text,
                    )
                    header_footer_count += 1

                footer_text = self._extract_container(
                    section.footer
                )

                if footer_text:
                    sections.append(
                        footer_text
                    )
                    header_footer_count += 1

        text = "\n\n".join(sections).strip()

        metadata = self._build_metadata(
            document=document,
            file_name=file_name,
            text=text,
            paragraph_count=paragraph_count,
            table_count=table_count,
            table_row_count=table_row_count,
            header_footer_count=header_footer_count,
        )

        metadata["text_extraction_status"] = (
            "success" if text else "empty"
        )

        return text, metadata

    def _extract_paragraph_text(
        self,
        paragraph: Any,
    ) -> str:
        """Extract and normalize a DOCX paragraph."""

        if self.include_hyperlinks:
            text = self._extract_paragraph_with_links(
                paragraph
            )
        else:
            text = paragraph.text or ""

        return self._normalize_text(text)

    def _extract_paragraph_with_links(
        self,
        paragraph: Any,
    ) -> str:
        """
        Extract paragraph text including hyperlink text.

        python-docx does not expose hyperlinks directly
        through paragraph.text, so XML elements are inspected.
        """

        try:
            from docx.oxml.ns import qn
        except ImportError:
            return paragraph.text or ""

        result: list[str] = []

        for child in paragraph._p:
            if child.tag == qn("w:hyperlink"):
                hyperlink_text = []

                for node in child.iter():
                    if node.tag == qn("w:t"):
                        if node.text:
                            hyperlink_text.append(
                                node.text
                            )

                if hyperlink_text:
                    result.append(
                        "".join(hyperlink_text)
                    )

            elif child.tag == qn("w:r"):
                for node in child.iter():
                    if node.tag == qn("w:t"):
                        if node.text:
                            result.append(node.text)

        extracted = "".join(result)

        if extracted:
            return extracted

        return paragraph.text or ""

    def _extract_table(
        self,
        table: Any,
    ) -> tuple[str, int]:
        """Extract table content row by row."""

        rows: list[str] = []

        for row in table.rows:
            cells: list[str] = []

            for cell in row.cells:
                cell_text = self._extract_cell_text(
                    cell
                )

                if cell_text:
                    cells.append(cell_text)

            if cells:
                rows.append(" | ".join(cells))

        return "\n".join(rows), len(table.rows)

    def _extract_cell_text(
        self,
        cell: Any,
    ) -> str:
        """Extract text from a table cell."""

        paragraphs: list[str] = []

        for paragraph in cell.paragraphs:
            text = self._extract_paragraph_text(
                paragraph
            )

            if text:
                paragraphs.append(text)

        return " ".join(paragraphs)

    def _extract_container(
        self,
        container: Any,
    ) -> str:
        """Extract paragraphs from a header/footer."""

        if container is None:
            return ""

        paragraphs: list[str] = []

        for paragraph in container.paragraphs:
            text = self._extract_paragraph_text(
                paragraph
            )

            if text:
                paragraphs.append(text)

        return "\n".join(paragraphs)

    @staticmethod
    def _normalize_text(text: str) -> str:
        """Normalize extracted DOCX text."""

        if not text:
            return ""

        text = text.replace("\r\n", "\n")
        text = text.replace("\r", "\n")
        text = text.replace("\x00", "")

        lines = []

        for line in text.split("\n"):
            line = " ".join(line.split())

            if line:
                lines.append(line)

        return "\n".join(lines)

    @staticmethod
    def _build_metadata(
        document: Any,
        file_name: str,
        text: str,
        paragraph_count: int,
        table_count: int,
        table_row_count: int,
        header_footer_count: int,
    ) -> dict[str, Any]:
        """Build metadata about the DOCX document."""

        core_properties = document.core_properties

        return {
            "source": "docx",
            "file_name": file_name,
            "character_count": len(text),
            "word_count": len(text.split()),
            "paragraph_count": paragraph_count,
            "table_count": table_count,
            "table_row_count": table_row_count,
            "header_footer_count": header_footer_count,
            "document_metadata": {
                "title": core_properties.title,
                "author": core_properties.author,
                "subject": core_properties.subject,
                "keywords": core_properties.keywords,
                "comments": core_properties.comments,
                "last_modified_by": (
                    core_properties.last_modified_by
                ),
                "created": str(
                    core_properties.created
                )
                if core_properties.created
                else None,
                "modified": str(
                    core_properties.modified
                )
                if core_properties.modified
                else None,
            },
        }

    @staticmethod
    def _looks_like_docx(data: bytes) -> bool:
        """
        Perform a lightweight DOCX signature check.

        DOCX files are ZIP containers and normally begin
        with the PK ZIP signature.
        """

        return data[:4] == b"PK\x03\x04"


def parse_docx(
    file_path: str | Path,
) -> tuple[str, dict[str, Any]]:
    """
    Convenience function for DOCX parsing.

    Example:
        text, metadata = parse_docx("resume.docx")
    """

    parser = DOCXParser()

    return parser.parse(file_path)